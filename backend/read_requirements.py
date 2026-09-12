from docx import Document

doc_path = r"c:\Users\kinal\OneDrive\Desktop\Neostat_project\AI_Engineer_Internship_Case_Study_Document_Intelligence_Final_Revised 2 (1).docx"
doc = Document(doc_path)

for para in doc.paragraphs:
    print(para.text)
    print("---")

print("\n\n=== TABLES ===")
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            print(cell.text, end=" | ")
        print()
    print("---")
